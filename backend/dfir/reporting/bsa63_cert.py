"""Draft certificate under Section 63 of the Bharatiya Sakshya Adhiniyam, 2023, for electronic records.

This is a DRAFT for the responsible person to review, correct and sign. Its wording is a
practical summary of what such certificates usually cover. It does not reproduce the
prescribed Schedule form, which must be followed exactly. A lawyer should confirm it before use."""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt

from ..core.config import host_info


def certificate_docx(case: dict, records: list[dict], package_sha256: str | None = None) -> bytes:
    d = Document()
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading("DRAFT — Certificate regarding electronic records", 0)
    p = d.add_paragraph()
    p.add_run("Section 63, Bharatiya Sakshya Adhiniyam, 2023 — DRAFT for review. ").bold = True
    p.add_run("Use the form prescribed in the Schedule to the Act. This draft only collects the technical facts the "
              "certificate needs. It must be reviewed by a lawyer, completed and signed by the person in charge of "
              "the device/system that produced these records, and (where required) by an expert.")
    d.add_heading("1. Identification of the electronic records", 1)
    t = d.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(("File", "Kind", "SHA-256", "MD5", "Acquired (UTC)")):
        t.rows[0].cells[i].text = h
    for r in records:
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text, c[3].text = r["filename"], r["kind"], r["sha256"], r["md5"]
        c[4].text = (r.get("stored_at") or {}).get("utc", "")
    if package_sha256:
        d.add_paragraph(f"Evidence package SHA-256: {package_sha256}")
    hi = host_info()
    d.add_heading("2. Device / system that produced the records", 1)
    d.add_paragraph(f"Computer: {hi['hostname']} — {hi['os']}")
    d.add_paragraph(f"Software: {hi['app']} (Python {hi['python']}); captures made with headless Chromium "
                    "(Playwright); hashes computed with SHA-256 and MD5.")
    d.add_paragraph("Owner/person in lawful control of the device: ______________________________")
    d.add_heading("3. Manner of production", 1)
    d.add_paragraph("Each record was produced by (a) direct upload of the file by the complainant, or (b) automated "
                    "capture of a public web page in a headless browser, or (c) HTTP download of media embedded in such "
                    "a page. Each record was hashed when it was acquired and stored write-once. Every action is logged "
                    "in a hash-chained custody log (custody_log.jsonl), which can be checked with verify_package.py.")
    d.add_heading("4. Statements to be confirmed by the signatory", 1)
    for s in ("The computer was operating properly during the period of production, or any malfunction did not affect "
              "the records.",
              "The records were produced in the ordinary course of the activities described above.",
              "The hash values above identify the records as produced."):
        d.add_paragraph(s, style="List Bullet")
    d.add_heading("5. Signature", 1)
    d.add_paragraph("Name: ____________________   Designation: ____________________")
    d.add_paragraph("Signature: ________________   Date: ____________   Place: ____________")
    d.add_paragraph(f"Case reference (local): {case.get('id')} — {case.get('title')}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
