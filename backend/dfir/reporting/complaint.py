"""Complaint preparation: NCRP-mapped field blocks + platform report texts.

The flow is Complaint Ready → Review/Edit → Submit, and the user always submits it themselves.
Nothing here contacts any authority or platform."""
from __future__ import annotations

from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path

import yaml
from sqlmodel import Session, select

from ..core.claims import sanitize_text
from ..core.custody import IST
from ..core.db import vault_for
from ..core.models import Case, Complaint, Location, Media
from ..pipeline import media_explanation
from ..reasoning.timeline import build_timeline

REFS = Path(__file__).parent / "refs"

INCIDENT_LABEL = {
    "ncii": "Non-consensual intimate imagery (morphed/AI-generated)",
    "sextortion": "Sextortion using manipulated media",
    "impersonation_fraud": "Impersonation / fraud using deepfake media",
    "voice_clone_scam": "Voice-clone scam",
    "political_defamation": "Manipulated media harming reputation (defamation/misinformation)",
    "other": "Deepfake / manipulated media incident",
}


@lru_cache
def legal_refs() -> dict:
    return yaml.safe_load((REFS / "legal_refs.yaml").read_text(encoding="utf-8"))


@lru_cache
def platform_routes() -> dict:
    return yaml.safe_load((REFS / "platform_routes.yaml").read_text(encoding="utf-8"))


def _lcfirst(s: str) -> str:
    return s[:1].lower() + s[1:]


def _ist(d: datetime | None) -> str:
    if not d:
        return "unknown"
    if d.tzinfo is None:
        d = d.replace(tzinfo=timezone.utc)
    return d.astimezone(IST).strftime("%d-%m-%Y %H:%M IST")


def draft_complaint(s: Session, case_id: str) -> dict:
    case = s.get(Case, case_id)
    media = s.exec(select(Media).where(Media.case_id == case_id)).all()
    locs = s.exec(select(Location).where(Location.case_id == case_id, Location.dismissed == False)).all()  # noqa: E712
    v = vault_for(case_id)
    primary = [m for m in media if m.role == "primary"]
    minors = case.safety.get("minors_possible")
    route = "women_child" if (minors or case.incident_type in ("ncii", "sextortion")) else "other"
    routes = platform_routes()

    dated = [l for l in locs if l.published_at or l.first_captured_at]
    first = min(dated, key=lambda l: l.published_at or l.first_captured_at) if dated else None
    first_time = (first.published_at or first.first_captured_at) if first else case.created_at

    findings = []
    for m in primary:
        ex = media_explanation(s, case_id, m.id)
        findings.append(f"{m.label}: {ex['headline']} (forensic indicators level: {ex['level']}).")
    live = [l for l in locs if l.state in ("live", "changed")]
    removed = [l for l in locs if l.state == "removed"]
    handles = sorted({(l.platform, l.account_handle) for l in locs if l.account_handle})
    plats = sorted({l.platform for l in locs}) or ["unknown"]

    short = (
        f"I am reporting {_lcfirst(INCIDENT_LABEL.get(case.incident_type, INCIDENT_LABEL['other']))}. "
        f"{case.description.strip()} "
        f"The media was first observed on {_ist(first_time)} on {first.platform if first else 'an online platform'}. "
        f"It has been found at {len(locs)} online location(s) ({len(live)} still live at last check). "
        + (f"Accounts observed posting it: {', '.join('@' + h + ' (' + p + ')' for p, h in handles[:6])}. " if handles else "")
        + "A forensic analysis tool assessed the media as follows: " + " ".join(findings) + " "
        "All URLs, screenshots and files have been preserved with SHA-256 hashes and a chain-of-custody log. "
        "I request removal of the content, preservation of account and upload data by the platforms, and investigation."
    )
    short = sanitize_text(" ".join(short.split()))
    if len(short) > 1500:
        short = short[:1490].rsplit(" ", 1)[0] + " …"

    tl = build_timeline(s, case_id)
    tl_lines = [f"- {(_ist(datetime.fromisoformat(e['at'])) if e['at'] and 'T' in e['at'] else (e['at'] or 'unknown'))}: "
                f"{e['title']} [{e['status']}; source: {e['time_source']}]" for e in tl["events"] if e["at"]][:40]

    evidence = []
    for rec in v.all_records():
        if rec["kind"] in ("original", "download", "capture_screenshot", "capture_html"):
            evidence.append({"file": rec["filename"], "kind": rec["kind"], "sha256": rec["sha256"], "md5": rec["md5"],
                             "size": rec["size"], "source": (rec.get("acquisition") or {}).get("page_url") or
                             (rec.get("acquisition") or {}).get("requested_url") or "uploaded by complainant",
                             "ncrp_uploadable": rec["size"] <= 10 * 1024 * 1024 and rec["kind"] in
                             ("capture_screenshot", "original")})
    long = "\n".join([
        f"INCIDENT: {INCIDENT_LABEL.get(case.incident_type)}",
        f"CASE REFERENCE (local): {case.id}",
        "",
        "DESCRIPTION", case.description or "(add a description in your own words)",
        "",
        "FORENSIC FINDINGS (indicators, not conclusions)", *[f"- {f}" for f in findings],
        "",
        "LOCATIONS WHERE THE MEDIA WAS FOUND",
        *[f"- {l.url} | {l.platform} | account: {('@' + l.account_handle) if l.account_handle else 'unknown'} | "
          f"published: {_ist(l.published_at)} | state at last check: {l.state} ({_ist(l.last_checked_at)})" for l in locs],
        "",
        "TIMELINE", *tl_lines,
        "",
        "WHAT IS NOT YET KNOWN",
        "- Who created the media and who operates the accounts listed above. The accounts are listed because they "
        "were observed posting the media. This is not an allegation that they created it.",
        "",
        "REQUESTED ACTION",
        "- Removal of the content from all platforms listed.",
        "- Preservation of account, login and upload records for the accounts listed, pending investigation.",
        "- Investigation into the creation and distribution of the media.",
    ])
    suspects = [{"type": "Social media account (observed posting)", "platform": p, "value": "@" + h,
                 "url": next((l.url for l in locs if l.account_handle == h and l.platform == p), "")} for p, h in handles]
    suspects += [{"type": "URL", "platform": l.platform, "value": l.url, "url": l.url} for l in locs]

    platform_reports = []
    for p in sorted({l.platform for l in locs}):
        r = routes["platforms"].get(p, routes["platforms"]["Website"])
        urls = [l.url for l in locs if l.platform == p and l.state != "removed"]
        if not urls:
            continue
        text = sanitize_text(
            f"I am the person depicted (or reporting on their behalf). The following content is "
            f"{'intimate/sexual imagery created without consent using AI manipulation' if case.incident_type in ('ncii', 'sextortion') else 'AI-manipulated media impersonating me'}"
            f". Please remove it and preserve related account data. URLs: {' '.join(urls)}. "
            f"Under India's IT Rules 2021, intermediaries are expected to act promptly on complaints about intimate or "
            f"impersonating (including artificially morphed) content.")
        platform_reports.append({"platform": p, "urls": urls, "how": r["report"], "help": r["help"],
                                 "grievance": r["grievance"], "text": text})

    provisions = [p for p in legal_refs()["provisions"]
                  if case.incident_type in p["incident_types"] or (minors and "minors" in p["incident_types"])]
    return {
        "route": {"key": route, "description": routes["ncrp"]["routes"][route], "portal": routes["ncrp"]["url"],
                  "helpline": routes["ncrp"]["helpline"]},
        "fields": {
            "incident_category": INCIDENT_LABEL.get(case.incident_type),
            "incident_datetime": _ist(first_time),
            "platforms": ", ".join(plats),
            "incident_description_short": short,
            "incident_description_long": long,
            "suspect_identifiers": suspects,
            "urls": [l.url for l in locs],
            "evidence_list": evidence,
            "complainant": case.complainant or {},
        },
        "platform_reports": platform_reports,
        "possibly_relevant_provisions": provisions,
        "legal_disclaimer": legal_refs()["disclaimer"],
        "notes": routes["ncrp"]["notes"],
        "other_routes": routes["other_routes"],
        "helplines": routes["helplines"],
    }


def get_or_create(s: Session, case_id: str, regenerate: bool = False) -> Complaint:
    c = s.get(Complaint, case_id)
    if c and c.state == "submitted" and not regenerate:
        return c
    draft = draft_complaint(s, case_id)
    if c is None:
        c = Complaint(case_id=case_id, fields=draft, state="ready")
    elif c.state == "ready" or regenerate:
        edited = (c.fields or {}).get("user_edits", {})
        draft["user_edits"] = edited
        c.fields = draft
    c.updated_at = datetime.now(timezone.utc)
    s.add(c)
    s.commit()
    return c


def complaint_docx(fields: dict, case_title: str) -> bytes:
    import io

    from docx import Document
    from docx.shared import Pt

    f = {**fields["fields"], **(fields.get("user_edits") or {})}
    d = Document()
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading(f"Cybercrime complaint — prepared draft ({case_title})", 0)
    d.add_paragraph("Prepared by the DFIR investigation tool for the complainant to review, edit and submit "
                    "themselves on https://cybercrime.gov.in. Not submitted to any authority.").italic = True
    for label, key in (("Incident category", "incident_category"), ("Date & time of incident", "incident_datetime"),
                       ("Platform(s)", "platforms")):
        d.add_paragraph().add_run(f"{label}: ").bold = True
        d.paragraphs[-1].add_run(str(f.get(key, "")))
    d.add_heading("Incident description (portal field, ≤1500 characters)", 1)
    d.add_paragraph(f.get("incident_description_short", ""))
    d.add_heading("Detailed statement", 1)
    for line in f.get("incident_description_long", "").splitlines():
        d.add_paragraph(line)
    d.add_heading("Accounts and URLs observed (not allegations of authorship)", 1)
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(("Type", "Platform", "Identifier")):
        t.rows[0].cells[i].text = h
    for sid in f.get("suspect_identifiers", []):
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text = sid["type"], sid["platform"], sid["value"]
    d.add_heading("Evidence list", 1)
    t = d.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(("File", "Kind", "SHA-256", "Source")):
        t.rows[0].cells[i].text = h
    for e in f.get("evidence_list", []):
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text, r[3].text = e["file"], e["kind"], e["sha256"], e["source"][:80]
    d.add_heading("Possibly relevant provisions (for your lawyer/police to confirm)", 1)
    d.add_paragraph(fields.get("legal_disclaimer", ""))
    for p in fields.get("possibly_relevant_provisions", []):
        d.add_paragraph(f"{p['law']} — {p['section']}: {p['title']}", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
